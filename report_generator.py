# -*- coding: utf-8 -*-
"""
风险评估报告生成器 - Report Generator
基于风险引擎结果，生成结构化Word格式的风险评估报告
"""

from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml
from datetime import datetime
import os

from risk_engine import AssessmentResult, RiskLevel


# ── 样式辅助函数 ──
def _set_cell_shading(cell, color_hex: str):
    """设置单元格底色"""
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)


def _add_run(paragraph, text: str, bold=False, size=10, color=None, font_name="微软雅黑"):
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    if color:
        run.font.color.rgb = RGBColor(*color)
    return run


def _risk_color(level: RiskLevel) -> str:
    return {"低风险": "6BCB77", "中等风险": "FFD93D", "高风险": "FF6B6B", "极高风险": "C00000"}.get(level.value, "808080")


def _risk_text_color(level: RiskLevel) -> tuple:
    return {"低风险": (0, 100, 0), "中等风险": (180, 130, 0), "高风险": (200, 0, 0), "极高风险": (120, 0, 0)}.get(level.value, (100, 100, 100))


# ── 报告生成主函数 ──
def generate_report(result: AssessmentResult, output_path: str = None) -> str:
    """生成完整的风险评估报告（Word格式）"""
    doc = Document()

    # ── 页面设置 ──
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    # ── 封面 ──
    for _ in range(6):
        doc.add_paragraph()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_run(title, "企业风险评估报告", bold=True, size=26, color=(31, 78, 121))

    doc.add_paragraph()

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_run(sub, result.company_name, bold=True, size=18, color=(46, 117, 182))

    doc.add_paragraph()
    doc.add_paragraph()

    info = doc.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_run(info, f"报告日期：{result.report_date}", size=12, color=(100, 100, 100))

    info2 = doc.add_paragraph()
    info2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_run(info2, "本报告基于企业风险管理信息搜集表自动生成", size=11, color=(130, 130, 130))

    info3 = doc.add_paragraph()
    info3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_run(info3, "仅供内部管理参考，请勿外传", size=10, color=(180, 180, 180))

    doc.add_page_break()

    # ── 目录页 ──
    toc_title = doc.add_heading("目  录", level=1)
    for run in toc_title.runs:
        run.font.color.rgb = RGBColor(31, 78, 121)

    toc_items = [
        "一、评估概要",
        "二、综合风险评分",
        "三、各维度风险分析",
        "    3.1 财务风险",
        "    3.2 负债与偿债风险",
        "    3.3 经营风险",
        "    3.4 行业与市场风险",
        "    3.5 法律与合规风险",
        "    3.6 税务风险",
        "    3.7 公司治理风险",
        "    3.8 供应链风险",
        "    3.9 技术与信息安全风险",
        "    3.10 人力资源风险",
        "    3.11 ESG与可持续发展风险",
        "    3.12 生产运营风险",
        "    3.13 安全生产风险",
        "    3.14 环境风险",
        "    3.15 信用风险",
        "    3.16 关联方与集团风险",
        "    3.17 项目投资风险",
        "四、风险矩阵与优先级排序",
        "五、核心风险总结",
        "六、附录：原始数据摘要",
    ]
    for item in toc_items:
        p = doc.add_paragraph()
        _add_run(p, item, size=12, color=(50, 50, 50))

    doc.add_page_break()

    # ── 一、评估概要 ──
    h1 = doc.add_heading("一、评估概要", level=1)
    for run in h1.runs:
        run.font.color.rgb = RGBColor(31, 78, 121)

    p = doc.add_paragraph()
    _add_run(p, "本报告对", size=11)
    _add_run(p, result.company_name, bold=True, size=11)
    _add_run(p, "的风险状况进行全面评估。评估基于企业风险管理信息搜集表所填数据，从财务、负债、行业市场、法律合规、税务、公司治理、供应链、技术信息、人力资源、ESG等17个维度进行系统性分析。", size=11)

    doc.add_paragraph()
    summary_table = doc.add_table(rows=4, cols=2, style="Table Grid")
    summary_table.alignment = WD_TABLE_ALIGNMENT.CENTER

    labels = ["企业名称", "综合风险评分", "综合风险等级", "评估日期"]
    values = [result.company_name, f"{result.overall_score:.2f} / 4.00", result.overall_level.value, result.report_date]

    for i, (label, val) in enumerate(zip(labels, values)):
        c0 = summary_table.cell(i, 0)
        c1 = summary_table.cell(i, 1)
        c0.width = Cm(4)
        c1.width = Cm(12)

        _add_run(c0.paragraphs[0], label, bold=True, size=11)
        _set_cell_shading(c0, "D6E4F0")

        if i == 2:
            _add_run(c1.paragraphs[0], val, bold=True, size=12, color=_risk_text_color(result.overall_level))
            _set_cell_shading(c1, _risk_color(result.overall_level))
        else:
            _add_run(c1.paragraphs[0], val, size=11)

    doc.add_page_break()

    # ── 二、综合风险评分 ──
    h2 = doc.add_heading("二、综合风险评分", level=1)
    for run in h2.runs:
        run.font.color.rgb = RGBColor(31, 78, 121)

    p = doc.add_paragraph()
    _add_run(p, f"综合风险评分：", size=12)
    _add_run(p, f"{result.overall_score:.2f}/4.00", bold=True, size=14, color=_risk_text_color(result.overall_level))
    _add_run(p, f"  |  等级：", size=12)
    _add_run(p, result.overall_level.value, bold=True, size=14, color=_risk_text_color(result.overall_level))

    doc.add_paragraph()

    # 评分表
    dim_table = doc.add_table(rows=len(result.dimensions) + 2, cols=5, style="Table Grid")
    dim_table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # 表头
    dim_headers = ["风险维度", "权重", "评分", "风险等级", "关键风险数"]
    for i, h in enumerate(dim_headers):
        cell = dim_table.cell(0, i)
        _add_run(cell.paragraphs[0], h, bold=True, size=10, color=(255, 255, 255))
        _set_cell_shading(cell, "2E75B6")

    # 数据行
    sorted_dims = sorted(result.dimensions.values(), key=lambda d: d.score, reverse=True)
    for r_idx, dim in enumerate(sorted_dims, 1):
        row_data = [dim.name, f"{dim.weight}%", f"{dim.score:.2f}", dim.level.value, str(len(dim.key_risks))]
        for c_idx, val in enumerate(row_data):
            cell = dim_table.cell(r_idx, c_idx)
            _add_run(cell.paragraphs[0], val, size=10)
            if c_idx == 3:
                _set_cell_shading(cell, _risk_color(dim.level))

    # 合计行
    total_row = len(result.dimensions) + 1
    dim_table.cell(total_row, 0).merge(dim_table.cell(total_row, 1))
    _add_run(dim_table.cell(total_row, 0).paragraphs[0], "综合评分", bold=True, size=11)
    _set_cell_shading(dim_table.cell(total_row, 0), "D6E4F0")
    _add_run(dim_table.cell(total_row, 2).paragraphs[0], f"{result.overall_score:.2f}", bold=True, size=11)
    _set_cell_shading(dim_table.cell(total_row, 2), "D6E4F0")
    _add_run(dim_table.cell(total_row, 3).paragraphs[0], result.overall_level.value, bold=True, size=11)
    _set_cell_shading(dim_table.cell(total_row, 3), _risk_color(result.overall_level))
    _add_run(dim_table.cell(total_row, 4).paragraphs[0], str(sum(len(d.key_risks) for d in result.dimensions.values())), bold=True, size=11)
    _set_cell_shading(dim_table.cell(total_row, 4), "D6E4F0")

    doc.add_page_break()

    # ── 三、各维度风险分析 ──
    h3 = doc.add_heading("三、各维度风险分析", level=1)
    for run in h3.runs:
        run.font.color.rgb = RGBColor(31, 78, 121)

    dim_names_cn = [
        "财务风险", "负债与偿债风险", "经营风险", "行业与市场风险",
        "法律与合规风险", "税务风险", "公司治理风险", "供应链风险",
        "技术与信息安全风险", "人力资源风险", "ESG与可持续发展风险",
        "生产运营风险", "安全生产风险", "环境风险", "信用风险",
        "关联方与集团风险", "项目投资风险",
    ]

    for idx, dim_name in enumerate(dim_names_cn, 1):
        dim = result.dimensions.get(dim_name)
        if not dim:
            continue

        # 维度标题
        h_dim = doc.add_heading(f"3.{idx}  {dim.name}", level=2)
        for run in h_dim.runs:
            run.font.color.rgb = RGBColor(46, 117, 182)

        # 评分摘要
        p = doc.add_paragraph()
        _add_run(p, "评分：", bold=True, size=11)
        _add_run(p, f"{dim.score:.2f}/4.00", bold=True, size=12, color=_risk_text_color(dim.level))
        _add_run(p, f"  风险等级：", bold=True, size=11)
        _add_run(p, dim.level.value, bold=True, size=12, color=_risk_text_color(dim.level))
        _add_run(p, f"  权重：{dim.weight}%", size=11)

        # 发现
        if dim.findings:
            p = doc.add_paragraph()
            _add_run(p, "主要发现：", bold=True, size=11)
            for finding in dim.findings:
                fp = doc.add_paragraph(style="List Bullet")
                _add_run(fp, finding, size=10)

        # 关键风险
        if dim.key_risks:
            p = doc.add_paragraph()
            _add_run(p, "关键风险点：", bold=True, size=11, color=(200, 0, 0))
            for kr in dim.key_risks:
                kp = doc.add_paragraph(style="List Bullet")
                _add_run(kp, kr, bold=True, size=10, color=(180, 0, 0))

        doc.add_paragraph()

    doc.add_page_break()

    # ── 四、风险矩阵与优先级排序 ──
    h4 = doc.add_heading("四、风险矩阵与优先级排序", level=1)
    for run in h4.runs:
        run.font.color.rgb = RGBColor(31, 78, 121)

    p = doc.add_paragraph()
    _add_run(p, "以下按风险评分从高到低排序，高评分维度应优先处理：", size=11)

    doc.add_paragraph()

    priority_table = doc.add_table(rows=len(sorted_dims) + 1, cols=4, style="Table Grid")
    priority_table.alignment = WD_TABLE_ALIGNMENT.CENTER

    pri_headers = ["优先级", "风险维度", "评分", "建议行动"]
    for i, h in enumerate(pri_headers):
        cell = priority_table.cell(0, i)
        _add_run(cell.paragraphs[0], h, bold=True, size=10, color=(255, 255, 255))
        _set_cell_shading(cell, "C55A11")

    for r_idx, dim in enumerate(sorted_dims, 1):
        priority = "★★★★★" if dim.score >= 3.5 else "★★★★" if dim.score >= 2.5 else "★★★" if dim.score >= 1.8 else "★★"
        action = _suggest_action(dim)
        row_data = [priority, dim.name, f"{dim.score:.2f}", action]
        for c_idx, val in enumerate(row_data):
            cell = priority_table.cell(r_idx, c_idx)
            _add_run(cell.paragraphs[0], val, size=10)
            if c_idx == 0:
                _set_cell_shading(cell, _risk_color(dim.level))

    doc.add_page_break()

    # ── 五、核心风险总结 ──
    h5 = doc.add_heading("五、核心风险总结", level=1)
    for run in h5.runs:
        run.font.color.rgb = RGBColor(31, 78, 121)

    all_key_risks = []
    for dim in sorted_dims:
        for kr in dim.key_risks:
            all_key_risks.append((dim.name, kr))

    if all_key_risks:
        p = doc.add_paragraph()
        _add_run(p, f"共识别 {len(all_key_risks)} 项关键风险点：", bold=True, size=11)

        kr_table = doc.add_table(rows=len(all_key_risks) + 1, cols=3, style="Table Grid")
        kr_table.alignment = WD_TABLE_ALIGNMENT.CENTER

        kr_headers = ["序号", "所属维度", "关键风险点"]
        for i, h in enumerate(kr_headers):
            cell = kr_table.cell(0, i)
            _add_run(cell.paragraphs[0], h, bold=True, size=10, color=(255, 255, 255))
            _set_cell_shading(cell, "C00000")

        for r_idx, (dim_name, kr) in enumerate(all_key_risks, 1):
            kr_table.cell(r_idx, 0).width = Cm(1.5)
            kr_table.cell(r_idx, 1).width = Cm(4)
            kr_table.cell(r_idx, 2).width = Cm(10)
            _add_run(kr_table.cell(r_idx, 0).paragraphs[0], str(r_idx), size=10)
            _add_run(kr_table.cell(r_idx, 1).paragraphs[0], dim_name, size=10)
            _add_run(kr_table.cell(r_idx, 2).paragraphs[0], kr, size=10)

    else:
        p = doc.add_paragraph()
        _add_run(p, "未识别出关键风险点，整体风险可控。", size=11)

    doc.add_page_break()

    # ── 六、附录：原始数据摘要 ──
    h6 = doc.add_heading("六、附录：原始数据摘要", level=1)
    for run in h6.runs:
        run.font.color.rgb = RGBColor(31, 78, 121)

    p = doc.add_paragraph()
    _add_run(p, "以下为各Sheet中提取的关键数据字段摘要：", size=11)
    doc.add_paragraph()

    for sheet_name, data in result.all_raw_data.items():
        if not data:
            continue
        h_sheet = doc.add_heading(f"  {sheet_name}", level=3)
        for run in h_sheet.runs:
            run.font.color.rgb = RGBColor(100, 100, 100)

        # 只显示有值的字段
        items = [(k, v) for k, v in data.items() if v is not None and str(v).strip()]
        if not items:
            p = doc.add_paragraph()
            _add_run(p, "（无数据）", size=10, color=(150, 150, 150))
            continue

        ncols = 2
        nrows = len(items) + 1
        data_table = doc.add_table(rows=nrows, cols=ncols, style="Table Grid")
        data_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        data_table.cell(0, 0).width = Cm(6)
        data_table.cell(0, 1).width = Cm(10)

        _add_run(data_table.cell(0, 0).paragraphs[0], "字段", bold=True, size=9, color=(255, 255, 255))
        _set_cell_shading(data_table.cell(0, 0), "808080")
        _add_run(data_table.cell(0, 1).paragraphs[0], "填写内容", bold=True, size=9, color=(255, 255, 255))
        _set_cell_shading(data_table.cell(0, 1), "808080")

        for r_idx, (k, v) in enumerate(items, 1):
            _add_run(data_table.cell(r_idx, 0).paragraphs[0], str(k), size=9)
            _add_run(data_table.cell(r_idx, 1).paragraphs[0], str(v), size=9)

        doc.add_paragraph()

    # ── 保存 ──
    if output_path is None:
        output_dir = r"D:\_Work\02_Documents\信息搜集表格"
        os.makedirs(output_dir, exist_ok=True)
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        output_path = os.path.join(output_dir, f"企业风险评估报告_{result.company_name}_{timestamp}.docx")

    doc.save(output_path)
    print(f"✓ 风险评估报告已生成: {output_path}")
    return output_path


def _suggest_action(dim) -> str:
 
