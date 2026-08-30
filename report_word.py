# -*- coding: utf-8 -*-
"""
Word 风险评估报告生成器 - Word Report Generator
生成专业格式的 Word (.docx) 风险评估报告
"""

import os
from datetime import datetime
from typing import Optional

from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

from report_base import BaseReportGenerator
from risk_engine import AssessmentResult, RiskLevel


# ── 样式辅助函数 ──
def _set_cell_shading(cell, color_hex: str):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)


def _add_run(paragraph, text: str, bold=False, size=10, color=None, font_name="微软雅黑"):
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    if color:
        run.font.color.rgb = color
    return run


def _set_cell_text(cell, text: str, bold=False, size=9, align=WD_ALIGN_PARAGRAPH.LEFT, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    _add_run(p, text, bold=bold, size=size, color=color)


def _set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}></w:tcBorders>')
    for edge, val in kwargs.items():
        element = parse_xml(
            f'<w:{edge} {nsdecls("w")} w:val="{val.get("val", "single")}" '
            f'w:sz="{val.get("sz", "4")}" w:color="{val.get("color", "000000")}" '
            f'w:space="0"/>'
        )
        tcBorders.append(element)
    tcPr.append(tcBorders)


class WordReportGenerator(BaseReportGenerator):
    """Word (.docx) 格式风险评估报告生成器"""

    # 颜色常量
    PRIMARY = RGBColor(0x1A, 0x3A, 0x5C)
    SECONDARY = RGBColor(0x2D, 0x6A, 0x9F)
    RED = RGBColor(0xC0, 0x00, 0x00)
    ORANGE = RGBColor(0xFF, 0x6B, 0x6B)
    YELLOW = RGBColor(0xFF, 0xD9, 0x3D)
    GREEN = RGBColor(0x6B, 0xCB, 0x77)
    GRAY = RGBColor(0x80, 0x80, 0x80)
    WHITE = RGBColor(0xFF, 0xFF, 0xFF)
    DARK_GRAY = RGBColor(0x66, 0x66, 0x66)
    LIGHT_BG = RGBColor(0xF5, 0xF7, 0xFA)

    def _risk_color_rgb(self, level_value: str) -> RGBColor:
        return {
            "低风险": self.GREEN,
            "中等风险": self.YELLOW,
            "高风险": self.ORANGE,
            "极高风险": self.RED,
        }.get(level_value, self.GRAY)

    def _add_heading_styled(self, doc, text: str, level=1):
        """添加带样式的标题"""
        heading = doc.add_heading(text, level=level)
        for run in heading.runs:
            run.font.color.rgb = self.PRIMARY if level == 1 else self.SECONDARY
            run.font.name = "微软雅黑"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        return heading

    def _add_body(self, doc, text: str, size=10, bold=False, color=None, space_after=Pt(6)):
        p = doc.add_paragraph()
        _add_run(p, text, bold=bold, size=size, color=color)
        p.paragraph_format.space_after = space_after
        p.paragraph_format.line_spacing = Pt(18)
        return p

    def _add_styled_table(self, doc, headers, rows, col_widths=None, header_bg="1A3A5C"):
        """添加带样式的表格"""
        table = doc.add_table(rows=1 + len(rows), cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = "Table Grid"

        # 表头
        for i, header in enumerate(headers):
            cell = table.rows[0].cells[i]
            _set_cell_text(cell, header, bold=True, size=9,
                           align=WD_ALIGN_PARAGRAPH.CENTER, color=self.WHITE)
            _set_cell_shading(cell, header_bg)

        # 数据行
        for r_idx, row in enumerate(rows):
            for c_idx, val in enumerate(row):
                cell = table.rows[r_idx + 1].cells[c_idx]
                _set_cell_text(cell, str(val), size=9)

                # 风险等级列（第3列）加背景色
                if c_idx == 2 and isinstance(val, str):
                    level_color = self._risk_color_rgb(val)
                    hex_str = str(level_color).lstrip("#")
                    _set_cell_shading(cell, hex_str)

        # 设置列宽
        if col_widths:
            for row in table.rows:
                for i, width in enumerate(col_widths):
                    row.cells[i].width = Cm(width)

        return table

    def generate(self) -> str:
        doc = Document()

        # ── 页面设置 ──
        section = doc.sections[0]
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

        sorted_dims = self._sorted_dimensions()
        key_risks = self._all_key_risks()

        # ── 封面 ──
        for _ in range(6):
            doc.add_paragraph()

        title = doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _add_run(title, "企业风险评估报告", bold=True, size=26, color=self.PRIMARY)

        doc.add_paragraph()

        company = doc.add_paragraph()
        company.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _add_run(company, self.result.company_name, bold=True, size=18, color=self.SECONDARY)

        for _ in range(4):
            doc.add_paragraph()

        meta_items = [
            f"报告日期：{self.result.report_date}",
            "本报告基于企业风险管理信息搜集表自动生成",
            "仅供内部管理参考，请勿外传",
        ]
        for item in meta_items:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _add_run(p, item, size=11, color=self.DARK_GRAY)

        doc.add_page_break()

        # ── 一、评估概要 ──
        self._add_heading_styled(doc, "一、评估概要", level=1)
        self._add_body(doc,
            f"本报告对 {self.result.company_name} 的风险状况进行全面评估。"
            "评估基于企业风险管理信息搜集表所填数据，从财务、负债、行业市场、法律合规、税务、"
            "公司治理、供应链、技术信息、人力资源、ESG等17个维度进行系统性分析。"
        )

        # 摘要表格
        summary_headers = ["项目", "内容"]
        summary_rows = [
            ["综合风险评分", f"{self.result.overall_score:.2f} / 4.00"],
            ["风险等级", self.result.overall_level.value],
            ["评估维度", f"{len(self.result.dimensions)} 个维度"],
            ["关键风险点", f"{len(key_risks)} 项待处理"],
        ]
        self._add_styled_table(doc, summary_headers, summary_rows,
                               col_widths=[5, 10], header_bg="2D6A9F")

        # 董事会级执行摘要（ISO 31000 / COSO 对齐）
        try:
            from risk_analytics import enrich_assessment
            basic = (self.result.all_raw_data or {}).get("企业基本信息", {})
            analytics = enrich_assessment(self.result, basic_info=basic)
            self._add_heading_styled(doc, "执行摘要（Executive Summary）", level=2)
            for line in analytics.get("executive_summary", []):
                self._add_body(doc, "• " + line)
            mat = analytics.get("erm_maturity", {})
            conf = analytics.get("confidence", {})
            self._add_body(doc,
                f"ERM 成熟度：{mat.get('level', '—')}/5（{mat.get('name', '')}）；"
                f"评估置信度：{conf.get('score', '—')}%（{conf.get('level', '')}）。"
                f" {analytics.get('standards_footer', '')}"
            )
            cross = analytics.get("cross_risk_alerts", [])
            if cross:
                self._add_body(doc, "交叉风险提示：" + "；".join(c["message"] for c in cross[:3]))
        except Exception:
            pass

        doc.add_paragraph()
        doc.add_page_break()

        # ── 二、各维度风险评分 ──
        self._add_heading_styled(doc, "二、各维度风险评分", level=1)
        self._add_body(doc, "以下按风险评分从高到低排序：")

        dim_headers = ["风险维度", "评分", "等级", "建议行动"]
        dim_rows = []
        for dim in sorted_dims:
            dim_rows.append([
                dim.name,
                f"{dim.score:.2f}",
                dim.level.value,
                self._suggest_action(dim),
            ])
        self._add_styled_table(doc, dim_headers, dim_rows,
                               col_widths=[4.5, 2, 2.5, 7])

        doc.add_paragraph()
        doc.add_page_break()

        # ── 三、各维度详细分析 ──
        self._add_heading_styled(doc, "三、各维度详细分析", level=1)

        for dim in sorted_dims:
            color = self._risk_color_rgb(dim.level.value)

            # 维度标题行
            p = doc.add_paragraph()
            _add_run(p, f"■ {dim.name}　", bold=True, size=11, color=color)
            _add_run(p, f"评分：{dim.score:.2f}　等级：{dim.level.value}",
                     size=10, color=color)
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)

            if dim.findings:
                for finding in dim.findings:
                    p = doc.add_paragraph()
                    _add_run(p, f"• {finding}", size=10)
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.left_indent = Cm(0.5)
            else:
                self._add_body(doc, "（无详细发现项）", size=10, color=self.GRAY)

            if dim.key_risks:
                p = doc.add_paragraph()
                _add_run(p, "关键风险：", bold=True, size=10, color=self.RED)
                _add_run(p, "；".join(dim.key_risks), size=10, color=self.RED)
                p.paragraph_format.space_after = Pt(4)

        doc.add_page_break()

        # ── 四、核心风险总结 ──
        self._add_heading_styled(doc, "四、核心风险总结", level=1)

        if key_risks:
            self._add_body(doc, f"共识别 {len(key_risks)} 项关键风险点：")

            kr_headers = ["序号", "所属维度", "关键风险点"]
            kr_rows = []
            for i, (dim_name, kr) in enumerate(key_risks, 1):
                kr_rows.append([str(i), dim_name, kr])

            self._add_styled_table(doc, kr_headers, kr_rows,
                                   col_widths=[1.5, 3.5, 11], header_bg="C00000")
        else:
            self._add_body(doc, "未识别出关键风险点，整体风险可控。")

        doc.add_paragraph()
        doc.add_page_break()

        # ── 五、风险登记册（ISO 31000） ──
        try:
            from risk_deep_analysis import run_deep_analysis
            from risk_analytics import compute_confidence
            basic = (self.result.all_raw_data or {}).get("企业基本信息", {})
            conf = compute_confidence(getattr(self.result, "form_stats", None)).get("score", 50)
            deep = run_deep_analysis(self.result, basic, confidence_pct=conf)
            self._add_heading_styled(doc, "五、风险登记册与应对策略", level=1)
            self._add_body(doc, deep.get("summary", ""))
            reg_headers = ["编号", "维度", "固有", "残余", "评级", "应对策略", "控制缺口"]
            reg_rows = []
            for r in deep.get("risk_register", [])[:12]:
                reg_rows.append([
                    r.get("id", ""), r.get("dimension", ""),
                    str(r.get("inherent_score", "")), str(r.get("residual_score", "")),
                    r.get("risk_rating", ""), r.get("treatment_primary", "")[:20],
                    r.get("control_gap", ""),
                ])
            if reg_rows:
                self._add_styled_table(doc, reg_headers, reg_rows,
                                       col_widths=[1.2, 2.5, 1, 1, 1, 2.5, 1.5], header_bg="1F4E79")
            for rec in deep.get("board_recommendations", [])[:5]:
                self._add_body(doc, f"• [{rec.get('priority')}] {rec.get('title')} — {rec.get('detail')}")
            mc = deep.get("monte_carlo")
            if mc:
                self._add_heading_styled(doc, "蒙特卡洛量化摘要", level=2)
                self._add_body(doc, mc.get("interpretation", ""))
            try:
                from data_external_risk import attach_external_evidence
                ext = attach_external_evidence(self.result)
                if ext and (ext.get("records") or ext.get("note")):
                    self._add_heading_styled(doc, "外部信用与司法证据", level=2)
                    self._add_body(doc, ext.get("note", "") + " " + (ext.get("disclaimer") or ""))
                    for rec in (ext.get("records") or [])[:6]:
                        self._add_body(
                            doc,
                            f"• [{rec.get('category_label')}] {rec.get('title')}（{rec.get('date') or '—'}）{rec.get('summary') or ''}"
                            + (f" 来源：{rec.get('url')}" if rec.get("url") else ""),
                        )
            except Exception:
                pass
            tld = deep.get("three_lines_of_defense")
            if tld:
                self._add_heading_styled(doc, "COSO 三道防线", level=2)
                self._add_body(doc, f"综合成熟度 {tld.get('overall_maturity')}，最弱防线：{tld.get('weakest_line')}")
            insights = deep.get("correlation_insights") or []
            if insights:
                self._add_heading_styled(doc, "风险关联传导", level=2)
                for ins in insights[:5]:
                    self._add_body(doc, f"• {ins.get('from')} → {ins.get('to')}：{ins.get('message')}（{ins.get('strength')}）")
            doc.add_page_break()
        except Exception:
            pass

        # ── 六、整改行动路线图（可执行） ──
        try:
            from action_planner import build_action_plans
            plans = build_action_plans(self.result)
            self._add_heading_styled(doc, "六、整改行动路线图", level=1)
            self._add_body(doc,
                "以下行动项按 ISO 31000 风险处理流程编制，含责任角色、时间节点与交付物。"
                f" 共 {plans.get('total_actions', 0)} 项，其中 P0 {plans.get('p0_count', 0)} 项须立即启动。"
            )
            act_headers = ["编号", "优先级", "维度", "行动", "责任人", "周期(天)"]
            act_rows = []
            for a in plans.get("action_items", [])[:15]:
                act_rows.append([
                    a.get("id", ""),
                    a.get("priority", ""),
                    a.get("dimension", ""),
                    a.get("title", "")[:40],
                    a.get("owner", ""),
                    str(a.get("timeline_days", "")),
                ])
            if act_rows:
                self._add_styled_table(doc, act_headers, act_rows,
                                       col_widths=[1.5, 1.2, 2.5, 5, 2.5, 1.5], header_bg="2D6A9F")
            doc.add_page_break()
        except Exception:
            pass

        # ── 七、数据闭环与 KRI 告警 ──
        try:
            from risk_closed_loop import build_export_closed_loop
            pkg = build_export_closed_loop(self.result)
            cl = pkg.get("closed_loop") or {}
            alerts = pkg.get("alerts") or {}
            v = cl.get("remediation_verification") or {}
            ts = cl.get("timeseries") or {}
            if cl or alerts.get("alerts"):
                self._add_heading_styled(doc, "七、数据闭环与 KRI 告警", level=1)
                if alerts.get("alerts"):
                    self._add_body(doc, f"告警摘要：{alerts.get('summary', '')}")
                    for a in alerts.get("alerts", [])[:8]:
                        self._add_body(doc, f"• [{a.get('severity', '').upper()}] {a.get('title')} — {a.get('message')}")
                if v:
                    self._add_body(doc,
                        f"整改验证：{v.get('verification_label', '—')}；"
                        f"评分变化 {v.get('overall_delta', '—')}；"
                        f"行动完成率 {v.get('actions_tracked', {}).get('completion_pct', 0)}%。"
                    )
                    self._add_body(doc, v.get("recommendation", ""))
                if ts.get("has_trend"):
                    self._add_body(doc, f"KPI 时序：{ts.get('message', '')}（{ts.get('period', '')}）")
                if cl.get("summary"):
                    self._add_body(doc, cl.get("summary"))
                doc.add_page_break()
        except Exception:
            pass

        # ── 附录：原始数据摘要 ──
        self._add_heading_styled(doc, "附录：原始数据摘要", level=1)
        self._add_body(doc, "以下为各Sheet中提取的关键数据字段摘要：")

        for sheet_name, data in self.result.all_raw_data.items():
            if not data:
                continue
            items = [(k, v) for k, v in data.items() if v is not None and str(v).strip()]
            if not items:
                continue

            self._add_heading_styled(doc, sheet_name, level=2)

            app_headers = ["字段", "填写内容"]
            app_rows = [(str(k), str(v)) for k, v in items]
            self._add_styled_table(doc, app_headers, app_rows,
                                   col_widths=[5, 11], header_bg="808080")
            doc.add_paragraph()

        # ── 保存 ──
        if self.output_path is None:
            output_dir = r"D:\_Work\02_Documents\信息搜集表格"
            os.makedirs(output_dir, exist_ok=True)
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            self.output_path = os.path.join(
                output_dir, f"企业风险评估报告_{self.result.company_name}_{timestamp}.docx"
            )

        doc.save(self.output_path)
        print(f"OK Word 风险评估报告已生成: {self.output_path}")
        return self.output_path
