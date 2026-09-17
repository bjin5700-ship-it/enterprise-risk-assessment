# -*- coding: utf-8 -*-
"""
企业风险动态评估系统 - Web 版
基于 Flask 的可视化风险评估平台
"""

import os
import sys
import re
import json
import uuid
from datetime import datetime, timedelta
from typing import Dict, Any, List, Tuple, Optional

from flask import Flask, render_template, request, jsonify, send_file, redirect, session
from werkzeug.utils import secure_filename

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
ROOT_DIR = os.path.abspath(os.path.join(BASE_DIR, os.pardir))
if ROOT_DIR not in sys.path:
    sys.path.insert(0, ROOT_DIR)

# 加载 web_app/.env（生产部署配置）
_env_file = os.path.join(BASE_DIR, ".env")
if os.path.isfile(_env_file):
    try:
        from dotenv import load_dotenv
        load_dotenv(_env_file)
    except ImportError:
        pass

from risk_engine import AssessmentResult, RiskLevel, SHEET_SCORERS, extract_all_data  # noqa: E402
from risk_analytics import enrich_assessment_full  # noqa: E402
from action_planner import build_action_plans  # noqa: E402

app = Flask(__name__)


@app.context_processor
def inject_erm_base():
    from flask import request
    from erm_auth import auth_enabled, current_user
    user = current_user()
    return {
        "erm_base": (request.script_root or "").rstrip("/"),
        "erm_user": user.to_dict() if user else None,
        "erm_auth_enabled": auth_enabled(),
    }

app.config["TEMPLATES_AUTO_RELOAD"] = True
app.config["MAX_CONTENT_LENGTH"] = 32 * 1024 * 1024  # 32MB

DATA_DIR = ROOT_DIR
EXPORT_DIR = os.path.join(BASE_DIR, "exports")
HISTORY_DIR = os.path.join(BASE_DIR, "data")
HISTORY_FILE = os.path.join(HISTORY_DIR, "history.json")
UPLOAD_DIR = os.path.join(BASE_DIR, "uploads")
BRIDGE_LATEST_FILE = os.path.join(HISTORY_DIR, "latest_assessment.json")
for _d in (EXPORT_DIR, HISTORY_DIR, UPLOAD_DIR):
    os.makedirs(_d, exist_ok=True)

_secret = (os.environ.get("ERM_SECRET_KEY") or "").strip()
_secret_file = os.path.join(HISTORY_DIR, ".secret")
if not _secret and os.path.isfile(_secret_file):
    try:
        _secret = open(_secret_file, "r", encoding="utf-8").read().strip()
    except OSError:
        _secret = ""
if not _secret:
    _secret = os.urandom(24).hex()
    try:
        with open(_secret_file, "w", encoding="utf-8") as _sf:
            _sf.write(_secret)
    except OSError:
        pass
app.secret_key = _secret or os.urandom(32)
app.config["SESSION_COOKIE_HTTPONLY"] = True
app.config["SESSION_COOKIE_SAMESITE"] = "Lax"
app.config["SESSION_COOKIE_NAME"] = "erm_session"
app.config["PERMANENT_SESSION_LIFETIME"] = timedelta(hours=12)
try:
    from server_config import is_production as _is_prod
    app.config["SESSION_COOKIE_SECURE"] = bool(_is_prod())
except Exception:
    app.config["SESSION_COOKIE_SECURE"] = False


@app.before_request
def _protect_routes():
    from erm_auth import enforce_request
    return enforce_request()

SHEET_NAME_MAP = {
    "企业基本信息": None,
    "经营风险": "经营风险",
    "财务风险": "财务风险",
    "生产运营风险": "生产运营风险",
    "安全生产风险": "安全生产风险",
    "环境风险": "环境风险",
    "法律合规风险": "法律与合规风险",
    "供应链风险": "供应链风险",
    "技术风险": "技术与信息安全风险",
    "人力资源风险": "人力资源风险",
    "信用风险": "信用风险",
    "综合评估": "综合评估",
    "关联方与集团风险": "关联方与集团风险",
    "项目投资风险": "项目投资风险",
    "行业政策风险": "行业与市场风险",
    "税务风险": "税务风险",
    "其他补充信息": None,
    "战略与声誉风险": "战略与声誉风险",
    "业务连续性风险": "业务连续性风险",
    "数据隐私合规风险": "数据隐私合规风险",
    "公司治理风险": "公司治理风险",
    "反贿赂道德合规风险": "反贿赂道德合规风险",
}

SCORING_SHEETS = {k for k, v in SHEET_NAME_MAP.items() if v is not None}
HISTORY_LIMIT = 100


def _normalize_field_type(raw: Any) -> str:
    s = str(raw or "文本").strip().lower()
    if s in ("dropdown", "select", "下拉"):
        return "下拉"
    if s in ("number", "numeric", "数值", "数字"):
        return "数值"
    if s in ("textarea", "longtext", "多行"):
        return "文本"
    if s in ("text", "string", "文本"):
        return "文本"
    return str(raw or "文本").strip() or "文本"


def _normalize_options(raw: Any) -> str:
    if raw is None:
        return ""
    text = str(raw).strip()
    if not text:
        return ""
    if "|" in text:
        return text
    if "," in text or "，" in text:
        parts = [p.strip() for p in re.split(r"[,，]", text) if p.strip()]
        return " | ".join(parts)
    return text


def _find_template_xlsx() -> Optional[str]:
    xlsx_files = [f for f in os.listdir(DATA_DIR) if f.endswith(".xlsx") and "企业风险" in f]
    if not xlsx_files:
        return None
    xlsx_files.sort(reverse=True)
    return os.path.join(DATA_DIR, xlsx_files[0])


def get_template_fields() -> Dict[str, Dict]:
    import openpyxl

    fp = _find_template_xlsx()
    if not fp:
        return {}

    wb = openpyxl.load_workbook(fp, data_only=True)
    template = {}
    for ws in wb.worksheets:
        sheet_name = ws.title
        fields = []
        for row in ws.iter_rows(min_row=5, max_row=ws.max_row, min_col=1, max_col=7):
            cells = [c.value for c in row]
            seq = cells[0]
            field_name = cells[1]
            field_type = cells[2]
            is_key = cells[3]
            options = cells[4]
            hint = cells[5] if len(cells) > 5 else None
            if isinstance(seq, (int, float)) and field_name:
                fields.append({
                    "seq": int(seq),
                    "name": str(field_name).strip(),
                    "type": _normalize_field_type(field_type),
                    "is_key": bool(is_key and "是" in str(is_key)),
                    "options": _normalize_options(options),
                    "hint": str(hint).strip() if hint else "",
                })
        if fields:
            template[sheet_name] = {"fields": fields, "scoring": sheet_name in SCORING_SHEETS}
    wb.close()
    from risk_matrix import inject_matrix_fields
    return inject_matrix_fields(template)


def normalize_form_data(raw: Dict[str, Dict[str, Any]], template: Dict[str, Dict]) -> Dict[str, Dict[str, Any]]:
    """将 Excel 或部分数据对齐到模板字段"""
    form_data: Dict[str, Dict[str, Any]] = {}
    for sheet_name, info in template.items():
        form_data[sheet_name] = {}
        imported = raw.get(sheet_name, {}) or {}
        for f in info.get("fields", []):
            val = imported.get(f["name"])
            if val is None:
                form_data[sheet_name][f["name"]] = ""
            elif isinstance(val, datetime):
                form_data[sheet_name][f["name"]] = val.strftime("%Y-%m-%d")
            else:
                form_data[sheet_name][f["name"]] = str(val).strip() if val != "" else ""
    return form_data


def import_excel_file(filepath: str) -> Dict[str, Dict[str, Any]]:
    template = get_template_fields()
    if not template:
        raise ValueError("Excel 模板未加载")
    raw = extract_all_data(filepath)
    return normalize_form_data(raw, template)


def compute_form_stats(form_data: Dict[str, Dict[str, Any]], template: Dict[str, Dict]) -> dict:
    total_fields = filled_fields = key_total = key_filled = 0
    sheet_stats = {}

    for sheet_name, info in template.items():
        fields = info.get("fields", [])
        sheet_total = len(fields)
        sheet_filled = 0
        for f in fields:
            total_fields += 1
            if f.get("is_key"):
                key_total += 1
            val = (form_data.get(sheet_name) or {}).get(f["name"], "")
            if val is not None and str(val).strip():
                filled_fields += 1
                sheet_filled += 1
                if f.get("is_key"):
                    key_filled += 1
        sheet_stats[sheet_name] = {
            "total": sheet_total,
            "filled": sheet_filled,
            "complete": sheet_filled >= sheet_total if sheet_total else False,
            "partial": 0 < sheet_filled < sheet_total,
        }

    scoring_sheets = [s for s, i in template.items() if i.get("scoring")]
    scoring_filled = sum(1 for s in scoring_sheets if sheet_stats.get(s, {}).get("filled", 0) > 0)

    return {
        "total_fields": total_fields,
        "filled_fields": filled_fields,
        "completion_pct": round(filled_fields / total_fields * 100, 1) if total_fields else 0,
        "key_total": key_total,
        "key_filled": key_filled,
        "key_completion_pct": round(key_filled / key_total * 100, 1) if key_total else 0,
        "sheet_stats": sheet_stats,
        "scoring_sheets_total": len(scoring_sheets),
        "scoring_sheets_filled": scoring_filled,
    }


def validate_form_data(form_data: Dict[str, Dict[str, Any]], template: Dict[str, Dict]) -> Tuple[List[str], List[str]]:
    errors: List[str] = []
    warnings: List[str] = []

    basic = form_data.get("企业基本信息", {})
    if not basic.get("企业名称") or not str(basic.get("企业名称", "")).strip():
        errors.append("请填写「企业基本信息」中的企业名称")
    uscc = str(basic.get("统一社会信用代码") or "").strip()
    if uscc:
        try:
            from credit_code import validate_credit_code
            ok, msg = validate_credit_code(uscc)
            if not ok:
                errors.append(msg)
        except Exception:
            pass

    stats = compute_form_stats(form_data, template)
    if stats["scoring_sheets_filled"] == 0:
        errors.append("请至少填写一个风险维度的数据后再评估")
    if stats["key_completion_pct"] < 30:
        warnings.append(f"关键字段完成度较低（{stats['key_filled']}/{stats['key_total']}），评估结果可能不够准确")
    if stats["completion_pct"] < 20:
        warnings.append("整体填写完成度较低，建议补充更多字段以获得可靠评估")
    # Phase-1 gate hints (soft; does not block unless strict)
    try:
        from risk_phase1 import build_data_quality_gate
        gate = build_data_quality_gate(stats, form_data, template)
        if gate.get("gate") == "block":
            warnings.append("【数据门禁】" + gate.get("gate_label", "数据不足"))
            for line in (gate.get("decision_impact") or [])[:2]:
                warnings.append(line)
        elif gate.get("gate") == "caution":
            warnings.append("【数据门禁】" + gate.get("gate_label", "建议补充数据"))
    except Exception:
        pass
    return errors, warnings


def run_assessment_from_form(form_data: Dict[str, Dict[str, Any]]) -> AssessmentResult:
    from risk_config import get_industry_profile

    dimensions = {}
    total_weighted = total_weight = 0.0

    company_name = "未命名企业"
    basic = form_data.get("企业基本信息", {})
    if basic.get("企业名称"):
        company_name = str(basic["企业名称"]).strip()

    industry_profile = get_industry_profile(basic.get("所属行业"))

    for excel_sheet, fields in form_data.items():
        scorer_key = SHEET_NAME_MAP.get(excel_sheet)
        if scorer_key is None:
            continue
        scorer = SHEET_SCORERS.get(scorer_key)
        if scorer is None:
            continue
        enriched = dict(fields)
        if excel_sheet in ("财务风险", "经营风险"):
            enriched["_industry_profile"] = industry_profile
        dim = scorer(enriched)
        dimensions[dim.name] = dim
        total_weighted += dim.score * dim.weight
        total_weight += dim.weight

    overall_score = round(total_weighted / total_weight, 2) if total_weight > 0 else 1.0
    return AssessmentResult(
        company_name=company_name,
        overall_score=overall_score,
        overall_level=RiskLevel.from_score(overall_score),
        dimensions=dimensions,
        all_raw_data=form_data,
        report_date=datetime.now().strftime("%Y年%m月%d日"),
    )


def assessment_to_dict(result: AssessmentResult, stats: dict = None, warnings: List[str] = None) -> dict:
    sorted_dims = sorted(result.dimensions.values(), key=lambda d: d.score, reverse=True)
    dims_data = [{
        "name": dim.name,
        "score": dim.score,
        "level": dim.level.value,
        "weight": dim.weight,
        "findings": dim.findings,
        "key_risks": dim.key_risks,
        "raw_data": dim.raw_data,
    } for dim in sorted_dims]

    all_key_risks = [
        {"dimension": dim.name, "risk": kr}
        for dim in result.dimensions.values()
        for kr in dim.key_risks
    ]

    level_counts = {"低风险": 0, "中等风险": 0, "高风险": 0, "极高风险": 0}
    for dim in result.dimensions.values():
        level_counts[dim.level.value] = level_counts.get(dim.level.value, 0) + 1

    basic = result.all_raw_data.get("企业基本信息", {}) if result.all_raw_data else {}
    from risk_bayesian import find_prior_assessment
    prior = find_prior_assessment(load_history(), result.company_name)

    payload = {
        "company_name": result.company_name,
        "overall_score": result.overall_score,
        "overall_level": result.overall_level.value,
        "report_date": result.report_date,
        "assessed_at": datetime.now().isoformat(),
        "dimensions": dims_data,
        "key_risks": all_key_risks,
        "all_raw_data": result.all_raw_data,
        "level_distribution": level_counts,
    }
    if stats is not None:
        payload["form_stats"] = stats
    if warnings:
        payload["warnings"] = warnings
    plans = build_action_plans(result)
    payload["action_plans"] = plans
    payload["analytics"] = enrich_assessment_full(
        result, stats, basic, payload.get("assessed_at"), prior_assessment=prior, action_plans=plans,
    )
    try:
        from risk_solution_program import build_solution_program

        deep = (payload.get("analytics") or {}).get("deep_analysis") or {}
        payload["solution_program"] = build_solution_program(
            result,
            action_plans=plans,
            deep_solutions=deep.get("deep_solutions") or payload.get("deep_solutions"),
            external_evidence=(payload.get("analytics") or {}).get("external_evidence"),
        )
    except Exception:
        payload["solution_program"] = None
    if payload["analytics"] and payload["analytics"].get("deep_analysis"):
        payload["deep_analysis"] = payload["analytics"]["deep_analysis"]
        payload["deep_solutions"] = payload["analytics"]["deep_analysis"].get("deep_solutions", [])
        payload["industry_playbook"] = payload["analytics"]["deep_analysis"].get("industry_playbook")
    if payload["analytics"] and payload["analytics"].get("bayesian_update"):
        payload["bayesian_update"] = payload["analytics"]["bayesian_update"]
    if payload["analytics"] and payload["analytics"].get("executive_brief"):
        payload["executive_brief"] = payload["analytics"]["executive_brief"]
    if payload["analytics"] and payload["analytics"].get("closed_loop"):
        payload["closed_loop"] = payload["analytics"]["closed_loop"]
        payload["alerts"] = payload["analytics"].get("alerts")
    if payload.get("analytics"):
        payload["phase1_enhancement"] = payload["analytics"].get("phase1_enhancement")
        payload["data_quality_gate"] = payload["analytics"].get("data_quality_gate")
        payload["verifiable_actions"] = payload["analytics"].get("verifiable_actions")
        payload["external_evidence"] = payload["analytics"].get("external_evidence")
    try:
        alerts_pkg = payload.get("alerts") or {}
        if alerts_pkg.get("alerts") or (payload.get("closed_loop") or {}).get("reassessment_trigger"):
            from risk_notifications import emit_assessment_alerts
            emit_assessment_alerts(
                payload["company_name"], alerts_pkg, payload.get("closed_loop"), source="assess",
            )
    except Exception:
        pass
    from risk_honesty import apply_honesty_gate
    apply_honesty_gate(payload)
    try:
        from risk_workflow import overlay_workflow
        overlay_workflow(payload)
    except Exception:
        pass
    try:
        _publish_latest_assessment(payload)
    except Exception:
        pass
    return payload


def _publish_latest_assessment(payload: dict) -> None:
    slim = {
        "company_name": payload.get("company_name"),
        "overall_score": payload.get("overall_score"),
        "overall_level": payload.get("overall_level"),
        "scoring_suppressed": payload.get("scoring_suppressed"),
        "data_quality_gate": payload.get("data_quality_gate"),
        "assessed_at": payload.get("assessed_at"),
        "form_stats": payload.get("form_stats"),
        "warnings": (payload.get("warnings") or [])[:8],
    }
    os.makedirs(HISTORY_DIR, exist_ok=True)
    with open(BRIDGE_LATEST_FILE, "w", encoding="utf-8") as f:
        json.dump(slim, f, ensure_ascii=False, indent=2)
    url = (os.environ.get("ERA_BRIDGE_NOTIFY_URL") or "").strip()
    if not url:
        return
    from urllib import request as urlrequest
    req = urlrequest.Request(
        url, data=json.dumps(slim, ensure_ascii=False).encode("utf-8"),
        headers={"Content-Type": "application/json"}, method="POST",
    )
    urlrequest.urlopen(req, timeout=3)


def generate_solution_data(result: AssessmentResult) -> dict:
    from solution_generator import SOLUTION_DB
    from risk_solution_program import build_solution_program

    high_dims = [dim for dim in result.dimensions.values()
                 if dim.level in (RiskLevel.HIGH, RiskLevel.CRITICAL) or dim.score >= 2.5]

    solutions = []
    priority_counts = {"P0": 0, "P1": 0, "P2": 0, "其他": 0}

    for dim in high_dims:
        sol = SOLUTION_DB.get(dim.name)
        if not sol:
            continue
        immediate = [{"measure": m, "priority": p, "detail": d} for m, p, d in sol.get("immediate", [])]
        short_term = [{"measure": m, "priority": p, "detail": d} for m, p, d in sol.get("short_term", [])]
        medium_term = [{"measure": m, "priority": p, "detail": d} for m, p, d in sol.get("medium_term", [])]
        for items in (immediate, short_term, medium_term):
            for item in items:
                p = item["priority"]
                if "P0" in p or p == "高":
                    priority_counts["P0"] += 1
                elif "P1" in p or p == "紧急":
                    priority_counts["P1"] += 1
                elif "P2" in p or p == "中":
                    priority_counts["P2"] += 1
                else:
                    priority_counts["其他"] += 1
        solutions.append({
            "dimension": dim.name,
            "score": dim.score,
            "level": dim.level.value,
            "title": sol.get("title", ""),
            "description": sol.get("description", ""),
            "immediate": immediate,
            "short_term": short_term,
            "medium_term": medium_term,
            "kpi": [{"name": k, "target": t, "freq": f} for k, t, f in sol.get("kpi", [])],
        })

    all_kpis = [
        {"dimension": dim.name, "name": kpi_name, "target": target, "freq": freq}
        for dim in high_dims
        for sol in [SOLUTION_DB.get(dim.name)]
        if sol and sol.get("kpi")
        for kpi_name, target, freq in sol["kpi"]
    ]

    out = {
        "company_name": result.company_name,
        "overall_score": result.overall_score,
        "overall_level": result.overall_level.value,
        "report_date": result.report_date,
        "solutions": solutions,
        "all_kpis": all_kpis,
        "high_risk_count": len(high_dims),
        "priority_counts": priority_counts,
        "total_measures": sum(priority_counts.values()),
        "action_plans": build_action_plans(result),
    }
    try:
        from risk_scenario import run_scenario_analysis
        from risk_kri import build_kri_dashboard
        out["scenario_analysis"] = run_scenario_analysis(result, result.all_raw_data)
        out["kri_dashboard"] = build_kri_dashboard(result, result.all_raw_data)
    except Exception:
        pass
    if result.all_raw_data:
        try:
            from risk_bayesian import find_prior_assessment
            template = get_template_fields()
            stats = compute_form_stats(result.all_raw_data, template) if template else None
            prior = find_prior_assessment(load_history(), result.company_name)
            plans = build_action_plans(result)
            analytics = enrich_assessment_full(
                result, stats, result.all_raw_data.get("企业基本信息", {}),
                prior_assessment=prior, action_plans=plans,
            )
            out["deep_analysis"] = analytics.get("deep_analysis")
            out["deep_solutions"] = (analytics.get("deep_analysis") or {}).get("deep_solutions", [])
            out["industry_playbook"] = (analytics.get("deep_analysis") or {}).get("industry_playbook")
            out["bayesian_update"] = analytics.get("bayesian_update")
            out["executive_brief"] = analytics.get("executive_brief")
            out["closed_loop"] = analytics.get("closed_loop")
            out["alerts"] = analytics.get("alerts")
            out["analytics"] = analytics
        except Exception:
            pass
    if out.get("analytics") and out["analytics"].get("data_quality_gate"):
        out["data_quality_gate"] = out["analytics"]["data_quality_gate"]
    try:
        out["solution_program"] = build_solution_program(
            result,
            action_plans=out.get("action_plans"),
            deep_solutions=out.get("deep_solutions"),
            external_evidence=(out.get("analytics") or {}).get("external_evidence"),
        )
    except Exception:
        out["solution_program"] = None
    from risk_honesty import apply_honesty_gate
    apply_honesty_gate(out)
    try:
        from risk_workflow import overlay_workflow
        overlay_workflow(out)
    except Exception:
        pass
    return out


def _safe_filename(name: str, ext: str, prefix: str = "企业风险评估报告") -> str:
    safe = re.sub(r'[\\/:*?"<>|]', "_", name).strip() or "未命名企业"
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    return f"{prefix}_{safe}_{ts}.{ext}"


# ── 评估档案 ──

def load_history() -> List[dict]:
    from erm_store import load_history as _load
    return _load()


def save_history(records: List[dict]) -> None:
    from erm_store import save_history as _save
    _save(records)


def add_history_record(assessment: dict, note: str = "") -> dict:
    from erm_store import add_history_record as _add
    return _add(assessment, note=note)


@app.route("/ping")
@app.route("/health")
@app.route("/api/health")
def ping():
    return jsonify({
        "ok": True,
        "status": "healthy",
        "service": "enterprise-risk-assessment",
        "version": "2.0.0-intl",
    })


@app.route("/api/engines")
def api_engines():
    try:
        from ocr_engine import ocr_status
        ocr = ocr_status()
    except Exception as exc:
        ocr = {"error": str(exc), "ready": False}
    try:
        import jsonschema  # noqa: F401
        js = True
    except Exception:
        js = False
    try:
        import pandera  # noqa: F401
        pa = True
    except Exception:
        pa = False
    try:
        import apscheduler  # noqa: F401
        aps = True
    except Exception:
        aps = False
    from erm_auth import admin_token, auth_enabled, current_user, ROLE_LABEL
    try:
        from data_external_risk import provider_status
        ext = provider_status()
    except Exception as exc:
        ext = {"provider": "none", "connected": False, "note": str(exc)}
    try:
        from erm_store import storage_status
        store = storage_status()
    except Exception as exc:
        store = {"history_backend": "unknown", "error": str(exc)}
    user = current_user()
    return jsonify({
        "ocr": ocr,
        "jsonschema": js,
        "pandera": pa,
        "apscheduler": aps,
        "admin_auth_required": bool(admin_token()),
        "auth_enabled": auth_enabled(),
        "auth": user.to_dict() if user else None,
        "roles": ROLE_LABEL,
        "external_risk": ext,
        "storage": store,
    })


@app.route("/api/validate/uscc")
def api_validate_uscc():
    from credit_code import validate_credit_code, normalize_credit_code
    code = request.args.get("code") or ""
    ok, msg = validate_credit_code(code)
    return jsonify({"ok": ok, "message": msg, "normalized": normalize_credit_code(code) if ok and code else ""})


@app.route("/api/external-risk", methods=["GET", "POST"])
def api_external_risk():
    try:
        from data_external_risk import lookup_external_risk
        if request.method == "POST":
            body = request.get_json(silent=True) or {}
            uscc = str(body.get("uscc") or body.get("code") or "")
            name = str(body.get("company_name") or body.get("name") or "")
        else:
            uscc = request.args.get("uscc") or request.args.get("code") or ""
            name = request.args.get("company_name") or request.args.get("name") or ""
        return jsonify(lookup_external_risk(uscc, name))
    except Exception as exc:
        return jsonify({"error": str(exc), "connected": False, "records": [], "suggested_fields": []}), 500


@app.route("/api/external-risk/demo-feed", methods=["GET"])
def api_external_risk_demo_feed():
    """HTTP 适配层：供 ERM_EXTERNAL_RISK_URL 指向的内置证据源。"""
    from data_external_risk import demo_feed_payload
    uscc = request.args.get("uscc") or request.args.get("code") or ""
    name = request.args.get("company_name") or request.args.get("name") or request.args.get("keyword") or ""
    return jsonify(demo_feed_payload(uscc, name))


def _safe_next(raw: str) -> str:
    s = str(raw or "/").strip()
    if s.startswith("/") and not s.startswith("//") and "://" not in s:
        return s
    return "/"


@app.route("/login", methods=["GET"])
def login_page():
    from erm_auth import auth_enabled, current_user
    user = current_user()
    if user:
        from erm_auth import app_path
        return redirect(app_path(_safe_next(request.args.get("next") or "/")))
    return render_template(
        "login.html",
        active_nav="",
        auth_enabled=auth_enabled(),
        next_url=_safe_next(request.args.get("next") or "/data-entry"),
    )


@app.route("/api/auth/me", methods=["GET"])
def api_auth_me():
    from erm_auth import auth_enabled, current_user
    user = current_user()
    return jsonify({
        "ok": True,
        "auth_enabled": auth_enabled(),
        "user": user.to_dict() if user else None,
    })


@app.route("/api/auth/login", methods=["POST"])
def api_auth_login():
    from erm_auth import authenticate, login_user
    from erm_rate_limit import check_allowed, client_key, record_failure, record_success

    body = request.get_json(silent=True) or {}
    username = str(body.get("username") or "")
    key = client_key(request.remote_addr or "", username)
    allowed, retry = check_allowed(key)
    if not allowed:
        return jsonify({
            "ok": False,
            "error": f"登录尝试过多，请 {retry} 秒后重试",
            "code": "ERR-ERM-RATE",
            "retry_after": retry,
        }), 429
    user = authenticate(username, str(body.get("password") or ""))
    if not user:
        record_failure(key)
        return jsonify({"ok": False, "error": "用户名或密码不正确", "code": "ERR-ERM-AUTH"}), 401
    record_success(key)
    login_user(user)
    return jsonify({"ok": True, "user": user.to_dict(), "next": _safe_next(body.get("next") or "/")})


@app.route("/api/auth/logout", methods=["GET", "POST"])
def api_auth_logout():
    from erm_auth import logout_user
    logout_user()
    if request.method == "GET" and "text/html" in (request.headers.get("Accept") or ""):
        from erm_auth import app_path
        return redirect(app_path("/"))
    return jsonify({"ok": True})


@app.route("/api/auth/session", methods=["POST"])
def api_auth_session():
    """兼容旧管理口令：等同于管理员登录。"""
    from erm_auth import admin_token, authenticate, login_user, ErmUser, default_admin_username
    body = request.get_json(silent=True) or {}
    got = str(body.get("token") or body.get("admin_token") or body.get("password") or "").strip()
    expected = admin_token()
    if expected and got and hmac_ok(got, expected):
        user = ErmUser(username=default_admin_username(), role="admin", source="token")
        login_user(user)
        session["erm_admin_token"] = got
        return jsonify({"ok": True, "required": True, "user": user.to_dict()})
    user = authenticate(str(body.get("username") or default_admin_username()), got)
    if user:
        login_user(user)
        return jsonify({"ok": True, "required": True, "user": user.to_dict()})
    if not expected:
        return jsonify({"ok": True, "required": False})
    return jsonify({"ok": False, "error": "口令不正确"}), 401


def hmac_ok(got: str, expected: str) -> bool:
    import hmac as _hmac
    if not got or not expected:
        return False
    return _hmac.compare_digest(got, expected)


@app.route("/api/bridge/latest")
def api_bridge_latest():
    if not os.path.isfile(BRIDGE_LATEST_FILE):
        return jsonify({"ok": False, "error": "尚无评估结果"}), 404
    with open(BRIDGE_LATEST_FILE, "r", encoding="utf-8") as f:
        return jsonify({"ok": True, "assessment": json.load(f)})


@app.route("/")
def index():
    return render_template("index.html", active_nav="dashboard")


@app.route("/data-entry")
def data_entry():
    return render_template("data_entry.html", active_nav="data-entry")


@app.route("/report")
def report():
    return render_template("report.html", active_nav="report")


@app.route("/solution")
def solution():
    return render_template("solution.html", active_nav="solution")


@app.route("/api/template")
def api_template():
    template = get_template_fields()
    if not template:
        return jsonify({"error": "未找到企业风险 Excel 模板，请将 xlsx 文件放在项目根目录"}), 404
    fp = _find_template_xlsx()
    return jsonify({"sheets": template, "template_file": os.path.basename(fp) if fp else None})


@app.route("/api/template/fields")
def api_template_fields():
    """兼容旧前端：仅返回 sheet 字段 dict"""
    template = get_template_fields()
    if not template:
        return jsonify({"error": "未找到 Excel 模板"}), 404
    return jsonify(template)


@app.route("/api/template/download")
def api_template_download():
    fp = _find_template_xlsx()
    if not fp:
        return jsonify({"error": "模板文件不存在"}), 404
    return send_file(fp, as_attachment=True, download_name=os.path.basename(fp))


@app.route("/api/import/excel", methods=["POST"])
def api_import_excel():
    try:
        if "file" not in request.files:
            return jsonify({"error": "请选择 Excel 文件"}), 400
        file = request.files["file"]
        if not file.filename:
            return jsonify({"error": "文件名为空"}), 400
        if not file.filename.lower().endswith((".xlsx", ".xlsm")):
            return jsonify({"error": "仅支持 .xlsx / .xlsm 格式"}), 400

        safe_name = secure_filename(file.filename) or "upload.xlsx"
        tmp_path = os.path.join(UPLOAD_DIR, f"{datetime.now().strftime('%Y%m%d%H%M%S')}_{safe_name}")
        file.save(tmp_path)
        try:
            form_data = import_excel_file(tmp_path)
        finally:
            if os.path.exists(tmp_path):
                os.remove(tmp_path)

        template = get_template_fields()
        stats = compute_form_stats(form_data, template)
        company = form_data.get("企业基本信息", {}).get("企业名称", "")
        return jsonify({
            "form_data": form_data,
            "form_stats": stats,
            "company_name": company or "未命名企业",
            "message": f"成功导入，共识别 {stats['filled_fields']} 个已填字段",
        })
    except Exception as exc:
        return jsonify({"error": f"导入失败: {exc}"}), 500



@app.route("/api/import/package", methods=["POST"])
def api_import_package():
    """Phase-1: 多文件包导入 — 主 Excel + 补充材料（PDF/Word/文本等）。"""
    try:
        from risk_phase1 import classify_supplement_file, extract_text_preview, build_supplement_package, build_data_quality_gate

        files = request.files.getlist("files") or []
        if not files and "file" in request.files:
            files = [request.files["file"]]
        if not files:
            return jsonify({"error": "请上传至少一个文件（主表 Excel + 可选补充材料）"}), 400

        main_form = None
        main_name = None
        supplements = []
        stamp = datetime.now().strftime("%Y%m%d%H%M%S")

        for f in files:
            if not f or not f.filename:
                continue
            raw_name = f.filename
            safe_name = secure_filename(raw_name) or "upload.bin"
            tmp_path = os.path.join(UPLOAD_DIR, f"{stamp}_{safe_name}")
            f.save(tmp_path)
            size = os.path.getsize(tmp_path) if os.path.exists(tmp_path) else 0
            meta = classify_supplement_file(raw_name, size)
            lower = raw_name.lower()
            try:
                if lower.endswith((".xlsx", ".xlsm")) and main_form is None:
                    main_form = import_excel_file(tmp_path)
                    main_name = raw_name
                    meta["role"] = "主数据"
                    meta["parsed"] = True
                else:
                    preview = extract_text_preview(tmp_path)
                    meta["parsed"] = bool(preview)
                    meta["text_preview"] = (preview[:400] + "…") if preview and len(preview) > 400 else preview
                    try:
                        from evidence_extract import suggest_fields_from_text
                        meta["suggested_fields"] = suggest_fields_from_text(preview or "", source=raw_name)
                    except Exception:
                        meta["suggested_fields"] = []
            finally:
                if os.path.exists(tmp_path):
                    os.remove(tmp_path)
            supplements.append(meta)

        if main_form is None:
            return jsonify({
                "error": "未识别到主搜集表 Excel（.xlsx/.xlsm）。请至少上传一份《企业风险信息搜集表》。",
                "supplements": build_supplement_package(supplements),
            }), 400

        template = get_template_fields()
        stats = compute_form_stats(main_form, template)
        pkg = build_supplement_package(supplements)
        suggested = []
        for meta in supplements:
            for item in meta.get("suggested_fields") or []:
                suggested.append(item)
        pkg["suggested_fields"] = suggested
        stats["_supplements"] = pkg
        gate = build_data_quality_gate(stats, main_form, template)
        company = (main_form.get("企业基本信息") or {}).get("企业名称", "") or "未命名企业"
        return jsonify({
            "form_data": main_form,
            "form_stats": {k: v for k, v in stats.items() if k != "_supplements"},
            "supplements": pkg,
            "suggested_fields": suggested,
            "data_quality_gate": gate,
            "company_name": company,
            "main_file": main_name,
            "message": f"已导入主表「{main_name}」+ 补充材料 {pkg.get('file_count', 0)} 份，识别字段 {stats.get('filled_fields', 0)} 个",
        })
    except Exception as exc:
        return jsonify({"error": f"文件包导入失败: {exc}"}), 500


@app.route("/api/import/local", methods=["GET"])
def api_import_local_list():
    """列出项目根目录下可导入的企业风险 Excel"""
    files = [
        {"name": f, "size": os.path.getsize(os.path.join(DATA_DIR, f))}
        for f in os.listdir(DATA_DIR)
        if f.endswith(".xlsx") and "企业风险" in f
    ]
    files.sort(key=lambda x: x["name"], reverse=True)
    return jsonify({"files": files})


@app.route("/api/import/local", methods=["POST"])
def api_import_local():
    try:
        filename = (request.get_json(force=True) or {}).get("filename")
        if not filename or ".." in filename:
            return jsonify({"error": "无效文件名"}), 400
        fp = os.path.join(DATA_DIR, filename)
        if not os.path.isfile(fp):
            return jsonify({"error": "文件不存在"}), 404
        form_data = import_excel_file(fp)
        template = get_template_fields()
        stats = compute_form_stats(form_data, template)
        return jsonify({
            "form_data": form_data,
            "form_stats": stats,
            "company_name": form_data.get("企业基本信息", {}).get("企业名称", "未命名企业"),
            "message": f"已从 {filename} 导入",
        })
    except Exception as exc:
        return jsonify({"error": f"导入失败: {exc}"}), 500


def _assess_form_payload(data: dict, *, strict: bool, save_history: bool, note: str) -> dict:
    template = get_template_fields()
    if not template:
        raise ValueError("Excel 模板未加载")

    errors, warnings = validate_form_data(data, template)
    if errors and strict:
        raise ValueError(errors[0])

    stats = compute_form_stats(data, template)
    result = run_assessment_from_form(data)
    payload = assessment_to_dict(result, stats=stats, warnings=warnings)
    if errors and not strict:
        payload["validation_errors"] = errors

    if save_history:
        record = add_history_record(payload, note=note)
        payload["history_id"] = record["id"]
    return payload


@app.route("/api/assess", methods=["POST"])
def api_assess():
    try:
        data = request.get_json(force=True) or {}
        payload = _assess_form_payload(
            data,
            strict=request.args.get("strict") == "1",
            save_history=request.args.get("save_history") == "1",
            note=request.args.get("note", ""),
        )
        return jsonify(payload)
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400
    except Exception as exc:
        return jsonify({"error": f"评估失败: {exc}"}), 500


@app.route("/api/assess/async", methods=["POST"])
def api_assess_async():
    """Queue full assessment; poll GET /api/assess/jobs/<id> or use webhook_url."""
    try:
        from erm_assess_jobs import create_job, run_job_async

        data = request.get_json(force=True) or {}
        webhook_url = str(data.pop("webhook_url", "") or request.args.get("webhook_url") or "").strip()
        webhook_url = webhook_url or (os.environ.get("ERA_BRIDGE_NOTIFY_URL") or "").strip()
        strict = request.args.get("strict") == "1"
        save_history = request.args.get("save_history") == "1"
        note = request.args.get("note", "")

        job_id = create_job()

        def _run() -> dict:
            return _assess_form_payload(data, strict=strict, save_history=save_history, note=note)

        run_job_async(job_id, _run, webhook_url=webhook_url)
        return jsonify(
            {
                "job_id": job_id,
                "status": "queued",
                "poll_url": f"/api/assess/jobs/{job_id}",
            }
        ), 202
    except Exception as exc:
        return jsonify({"error": f"异步评估入队失败: {exc}"}), 500


@app.route("/api/assess/jobs/<job_id>", methods=["GET"])
def api_assess_job(job_id: str):
    from erm_assess_jobs import get_job

    row = get_job(job_id)
    if not row:
        return jsonify({"error": "任务不存在"}), 404
    out = {
        "job_id": row["id"],
        "status": row["status"],
        "created_at": row["created_at"],
        "updated_at": row["updated_at"],
    }
    if row.get("error"):
        out["error"] = row["error"]
    if row.get("result") is not None:
        out["result"] = row["result"]
    return jsonify(out)


@app.route("/api/solution", methods=["POST"])
def api_solution():
    try:
        data = request.get_json(force=True) or {}
        template = get_template_fields()
        if template:
            data = normalize_form_data(data.get("form_data", data), template)
        result = run_assessment_from_form(data)
        return jsonify(generate_solution_data(result))
    except Exception as exc:
        return jsonify({"error": f"生成解决方案失败: {exc}"}), 500


@app.route("/api/history", methods=["GET"])
def api_history_list():
    from erm_auth import current_user, filter_history, is_demo_record
    all_records = load_history()
    records = filter_history(all_records, current_user())
    # 演示档案始终置顶，便于公网审阅与国际方法论展示
    demos = [r for r in all_records if is_demo_record(r)]
    seen = {str(r.get("id")) for r in records}
    for d in demos:
        if str(d.get("id")) not in seen:
            records.insert(0, d)
            seen.add(str(d.get("id")))
    summary = [{
        "id": r["id"],
        "company_name": r.get("company_name"),
        "overall_score": r.get("overall_score"),
        "overall_level": r.get("overall_level"),
        "assessed_at": r.get("assessed_at"),
        "note": r.get("note", ""),
    } for r in records]
    return jsonify({"records": summary, "total": len(summary)})


@app.route("/api/history", methods=["POST"])
def api_history_save():
    try:
        body = request.get_json(force=True) or {}
        assessment = body.get("assessment")
        if not assessment:
            return jsonify({"error": "缺少 assessment 数据"}), 400
        record = add_history_record(assessment, note=body.get("note", ""))
        return jsonify(record)
    except Exception as exc:
        return jsonify({"error": str(exc)}), 500


@app.route("/api/history/<record_id>", methods=["GET"])
def api_history_get(record_id):
    from erm_auth import current_user, can_read_record
    from erm_store import get_assessment
    rec = get_assessment(record_id)
    if not rec:
        return jsonify({"error": "档案不存在"}), 404
    if not can_read_record(rec, current_user()):
        return jsonify({"error": "未登录或无权查看该档案", "code": "ERR-ERM-AUTH"}), 401
    try:
        from risk_workflow import overlay_workflow
        if rec.get("assessment"):
            overlay_workflow(rec["assessment"])
    except Exception:
        pass
    return jsonify(rec)


@app.route("/api/history/<record_id>", methods=["DELETE"])
def api_history_delete(record_id):
    from erm_store import delete_assessment
    ok = delete_assessment(record_id)
    return jsonify({"ok": True, "deleted": ok})


@app.route("/api/export/<fmt>", methods=["POST"])
def api_export(fmt):
    allowed = {"docx", "html", "md", "pdf", "solution-docx"}
    if fmt not in allowed:
        return jsonify({"error": f"不支持的格式: {fmt}"}), 400
    try:
        data = request.get_json(force=True) or {}
        result = run_assessment_from_form(data)

        if fmt == "solution-docx":
            from solution_generator import generate_solution_report
            filename = _safe_filename(result.company_name, "docx", prefix="企业风险解决方案报告")
            out_path = os.path.join(EXPORT_DIR, filename)
            path = generate_solution_report(result, out_path)
            mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        elif fmt == "docx":
            from report_word import WordReportGenerator
            out_path = os.path.join(EXPORT_DIR, _safe_filename(result.company_name, "docx"))
            path = WordReportGenerator(result, out_path).generate()
            mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        elif fmt == "pdf":
            from report_pdf import PDFReportGenerator
            out_path = os.path.join(EXPORT_DIR, _safe_filename(result.company_name, "pdf"))
            path = PDFReportGenerator(result, out_path).generate()
            mime = "application/pdf"
        elif fmt == "html":
            from report_html import HTMLReportGenerator
            out_path = os.path.join(EXPORT_DIR, _safe_filename(result.company_name, "html"))
            path = HTMLReportGenerator(result, out_path).generate()
            mime = "text/html; charset=utf-8"
        else:
            from report_markdown import MarkdownReportGenerator
            out_path = os.path.join(EXPORT_DIR, _safe_filename(result.company_name, "md"))
            path = MarkdownReportGenerator(result, out_path).generate()
            mime = "text/markdown; charset=utf-8"

        try:
            from erm_objects import put_file
            put_file(
                path,
                kind="export",
                content_type=mime.split(";")[0].strip(),
                company_name=result.company_name,
            )
        except Exception:
            pass

        return send_file(path, as_attachment=True, download_name=os.path.basename(path), mimetype=mime)
    except Exception as exc:
        return jsonify({"error": f"导出失败: {exc}"}), 500


@app.route("/api/objects/<oid>", methods=["GET"])
def api_object_get(oid):
    try:
        from erm_store import get_evidence
        from erm_objects import open_object
        rec = get_evidence(oid)
        if not rec:
            return jsonify({"error": "对象不存在"}), 404
        path, data = open_object(rec)
        name = rec.get("filename") or "download"
        mime = rec.get("content_type") or "application/octet-stream"
        if data is not None:
            from io import BytesIO
            return send_file(BytesIO(data), as_attachment=True, download_name=name, mimetype=mime)
        return send_file(path, as_attachment=True, download_name=name, mimetype=mime)
    except FileNotFoundError:
        return jsonify({"error": "文件不存在"}), 404
    except Exception as exc:
        return jsonify({"error": str(exc)}), 500


# ── Phase D: 外部数据集成 ──

@app.route("/api/integrations", methods=["GET"])
def api_integrations_list():
    from data_integration import list_connectors
    return jsonify({"connectors": list_connectors()})


@app.route("/api/integrations/erp", methods=["POST"])
def api_integrations_erp():
    try:
        from data_integration import apply_erp_payload, merge_external_into_form, ERP_FIELD_MAP, save_integration_log
        payload = request.get_json(force=True) or {}
        from schema_validate import validate_erp_payload
        errs, warns = validate_erp_payload(payload)
        if errs:
            return jsonify({"error": errs[0], "errors": errs, "warnings": warns}), 400
        external = apply_erp_payload(payload)
        template = get_template_fields()
        base_form = normalize_form_data({}, template) if template else {}
        merged = merge_external_into_form(base_form, {**external, "_force": payload.get("force", False)}, ERP_FIELD_MAP)
        save_integration_log("erp", payload, {"fields_updated": len(external)})
        return jsonify({"form_data": merged, "updated_fields": external, "message": f"已同步 {len(external)} 个 ERP 字段"})
    except Exception as exc:
        return jsonify({"error": str(exc)}), 500


@app.route("/api/integrations/financial", methods=["POST"])
def api_integrations_financial():
    try:
        from data_integration import apply_financial_api_payload, merge_external_into_form, FINANCIAL_API_MAP, save_integration_log
        body = request.get_json(force=True) or {}
        external = apply_financial_api_payload(body)
        template = get_template_fields()
        base_form = normalize_form_data(body.get("form_data", {}), template) if template else {}
        merged = merge_external_into_form(base_form, {**external, "_force": body.get("force", False)}, FINANCIAL_API_MAP)
        save_integration_log("financial_api", body, {"fields_updated": len(external)})
        return jsonify({"form_data": merged, "updated_fields": external})
    except Exception as exc:
        return jsonify({"error": str(exc)}), 500


@app.route("/api/integrations/csv", methods=["POST"])
def api_integrations_csv():
    try:
        from data_integration import parse_csv_financial, merge_external_into_form, ERP_FIELD_MAP, save_integration_log
        content = request.get_data(as_text=True)
        if not content and "file" in request.files:
            content = request.files["file"].read().decode("utf-8-sig", errors="ignore")
        if not content:
            return jsonify({"error": "请提供 CSV 内容或文件"}), 400
        external = parse_csv_financial(content)
        body = request.form.to_dict() if request.files else (request.get_json(silent=True) or {})
        template = get_template_fields()
        base = body.get("form_data") or {}
        base_form = normalize_form_data(base, template) if template else {}
        merged = merge_external_into_form(base_form, external, ERP_FIELD_MAP)
        save_integration_log("csv", {"rows": len(external)}, {"fields_updated": len(external)})
        return jsonify({"form_data": merged, "updated_fields": external})
    except Exception as exc:
        return jsonify({"error": str(exc)}), 500


@app.route("/api/monitoring/snapshot", methods=["POST"])
def api_monitoring_snapshot():
    """合并外部数据 + 最近评估 → 实时 KRI 快照"""
    try:
        data = request.get_json(force=True) or {}
        template = get_template_fields()
        if not template:
            return jsonify({"error": "模板未加载"}), 404
        form_data = normalize_form_data(data.get("form_data", data), template)
        result = run_assessment_from_form(form_data)
        from risk_kri import build_kri_dashboard
        from risk_scenario import run_scenario_analysis
        kri = build_kri_dashboard(result, form_data)
        scenario = run_scenario_analysis(result, form_data)
        return jsonify({
            "company_name": result.company_name,
            "overall_score": result.overall_score,
            "overall_level": result.overall_level.value,
            "kri_dashboard": kri,
            "liquidity_status": scenario.get("liquidity_stress", {}).get("status"),
            "snapshot_at": datetime.now().isoformat(),
        })
    except Exception as exc:
        return jsonify({"error": str(exc)}), 500


@app.route("/api/scenario", methods=["POST"])
def api_scenario():
    try:
        data = request.get_json(force=True) or {}
        result = run_assessment_from_form(data)
        from risk_scenario import run_scenario_analysis
        return jsonify(run_scenario_analysis(result, data))
    except Exception as exc:
        return jsonify({"error": str(exc)}), 500


@app.route("/api/quant", methods=["POST"])
def api_quant():
    try:
        data = request.get_json(force=True) or {}
        template = get_template_fields()
        if template:
            data = normalize_form_data(data.get("form_data", data), template)
        result = run_assessment_from_form(data)
        from risk_quant import run_monte_carlo
        from risk_analytics import enrich_assessment_full
        basic = data.get("企业基本信息", {})
        stats = compute_form_stats(data, template) if template else None
        analytics = enrich_assessment_full(result, stats, basic)
        conf = analytics.get("confidence", {}).get("score", 50)
        appetite = 2.5
        for k, v in [("保守", 1.8), ("稳健", 2.2), ("平衡", 2.6), ("积极", 3.0)]:
            if k in str(basic.get("风险承受度", "")):
                appetite = v
                break
        return jsonify(run_monte_carlo(
            result, confidence_pct=conf, appetite_threshold=appetite, basic=basic,
        ))
    except Exception as exc:
        return jsonify({"error": str(exc)}), 500


@app.route("/api/alerts", methods=["GET"])
def api_alerts():
    company = request.args.get("company", "")
    records = load_history()
    if company:
        for r in records:
            if r.get("company_name") == company and r.get("assessment"):
                a = r["assessment"].get("analytics") or {}
                return jsonify(a.get("alerts") or {"alerts": [], "summary": "无缓存告警"})
    return jsonify({"alerts": [], "summary": "请指定 company 参数或先完成评估"})


@app.route("/api/kpi/timeseries", methods=["GET"])
def api_kpi_timeseries():
    company = request.args.get("company", "")
    if not company:
        return jsonify({"error": "缺少 company 参数"}), 400
    from risk_timeseries import get_company_timeseries, analyze_timeseries
    return jsonify({
        "timeseries": get_company_timeseries(company),
        "analysis": analyze_timeseries(company),
    })


@app.route("/api/notifications", methods=["GET"])
def api_notifications_list():
    from risk_notifications import list_notifications
    unread = request.args.get("unread") == "1"
    company = request.args.get("company", "")
    return jsonify(list_notifications(unread_only=unread, company=company or None))


@app.route("/api/notifications/read", methods=["POST"])
def api_notifications_read():
    from risk_notifications import mark_notifications_read
    body = request.get_json(force=True) or {}
    return jsonify(mark_notifications_read(ids=body.get("ids"), mark_all=body.get("mark_all", False)))


@app.route("/api/notifications/config", methods=["GET", "PUT"])
def api_notifications_config():
    from risk_notifications import load_notification_config, save_notification_config
    if request.method == "GET":
        cfg = load_notification_config()
        safe = dict(cfg)
        if safe.get("smtp_password"):
            safe["smtp_password"] = "******"
        return jsonify(safe)
    body = request.get_json(force=True) or {}
    for k in ("smtp_password", "smtp_host", "smtp_port", "smtp_user", "smtp_use_tls"):
        body.pop(k, None)
    cfg = load_notification_config()
    cfg.update(body)
    save_notification_config(cfg)
    return jsonify({"ok": True})


@app.route("/api/integrations/schedule", methods=["GET", "PUT"])
def api_integrations_schedule():
    from data_integration import load_sync_schedule, save_sync_schedule, normalize_schedule
    from sync_scheduler import scheduler_status
    if request.method == "GET":
        cfg = load_sync_schedule()
        return jsonify({"schedule": cfg, "scheduler": scheduler_status()})
    body = request.get_json(force=True) or {}
    incoming = body.get("schedule") or body
    cfg = load_sync_schedule()
    if "jobs" in incoming:
        cfg = normalize_schedule({**cfg, **incoming})
    else:
        cfg.update({k: v for k, v in incoming.items() if k in ("global_enabled", "max_parallel_jobs")})
    save_sync_schedule(cfg)
    return jsonify({"ok": True, "schedule": load_sync_schedule()})


@app.route("/api/integrations/schedule/jobs", methods=["POST"])
def api_schedule_upsert_job():
    from data_integration import upsert_sync_job
    body = request.get_json(force=True) or {}
    job = upsert_sync_job(body.get("job") or body)
    return jsonify({"ok": True, "job": job})


@app.route("/api/integrations/schedule/jobs/<job_id>", methods=["DELETE"])
def api_schedule_delete_job(job_id):
    from data_integration import delete_sync_job
    return jsonify({"ok": delete_sync_job(job_id)})


@app.route("/api/integrations/schedule/run", methods=["POST"])
def api_integrations_schedule_run():
    from data_integration import execute_due_jobs_parallel, execute_scheduled_sync, load_sync_schedule
    job_id = request.args.get("job_id")
    if job_id:
        sched = load_sync_schedule()
        job = next((j for j in sched.get("jobs", []) if j.get("id") == job_id), None)
        if not job:
            return jsonify({"error": f"任务不存在: {job_id}"}), 404
        return jsonify(execute_scheduled_sync(job, force=True))
    job_ids = request.get_json(silent=True) or {}
    ids = job_ids.get("job_ids") if isinstance(job_ids, dict) else None
    return jsonify(execute_due_jobs_parallel(force=True, job_ids=ids))


@app.route("/api/integrations/oauth/providers", methods=["GET"])
def api_oauth_providers():
    from erp_oauth import load_oauth_providers, oauth_connection_status
    doc = load_oauth_providers()
    safe = []
    for p in doc.get("providers", []):
        sp = dict(p)
        if sp.get("client_secret"):
            sp["client_secret"] = "******"
        safe.append(sp)
    return jsonify({"providers": safe, "connections": oauth_connection_status()})


@app.route("/api/integrations/oauth/providers/<provider_id>", methods=["PUT"])
def api_oauth_provider_update(provider_id):
    from erp_oauth import update_provider
    body = request.get_json(force=True) or {}
    body.pop("client_secret", None)
    provider = update_provider(provider_id, body)
    if provider.get("client_secret"):
        provider = dict(provider)
        provider["client_secret"] = "******"
    return jsonify({"ok": True, "provider": provider})


@app.route("/api/integrations/oauth/authorize", methods=["GET"])
def api_oauth_authorize():
    from erp_oauth import build_authorization_url
    from server_config import get_oauth_redirect_uri
    provider_id = request.args.get("provider_id", "generic")
    company = request.args.get("company", "")
    redirect_uri = request.args.get("redirect_uri") or get_oauth_redirect_uri()
    return jsonify(build_authorization_url(provider_id, redirect_uri, company))


@app.route("/api/integrations/oauth/callback")
def api_oauth_callback():
    from erp_oauth import handle_oauth_callback
    from erm_auth import app_path
    err = request.args.get("error")
    if err:
        return redirect(f"{app_path('/data-entry')}?oauth_error={err}")
    try:
        code = request.args.get("code", "")
        state = request.args.get("state", "")
        result = handle_oauth_callback(code, state)
        return redirect(f"{app_path('/data-entry')}?oauth_ok={result.get('provider_id', '')}")
    except Exception as exc:
        return redirect(f"{app_path('/data-entry')}?oauth_error={exc}")


@app.route("/api/integrations/oauth/sync", methods=["POST"])
def api_oauth_sync():
    try:
        body = request.get_json(force=True) or {}
        provider_id = body.get("provider_id", "generic")
        company = body.get("company_name", "")
        from erp_oauth import fetch_oauth_erp_payload
        from data_integration import merge_external_into_form, ERP_FIELD_MAP, FINANCIAL_API_MAP

        fetched = fetch_oauth_erp_payload(provider_id)
        external = fetched.get("external") or {}
        template = get_template_fields()
        if not template:
            return jsonify({"error": "模板未加载"}), 404
        base = normalize_form_data(body.get("form_data", {}), template)
        field_map = ERP_FIELD_MAP if fetched.get("type") == "oauth_erp" else FINANCIAL_API_MAP
        merged = merge_external_into_form(base, {**external, "_force": True}, field_map)
        if company:
            merged.setdefault("企业基本信息", {})["企业名称"] = company
        return jsonify({"form_data": merged, "updated_fields": external, "message": f"OAuth 已同步 {len(external)} 个字段"})
    except Exception as exc:
        return jsonify({"error": str(exc)}), 500


@app.route("/api/notifications/devices", methods=["GET", "POST", "DELETE"])
def api_mobile_devices():
    from risk_mobile_push import register_device, unregister_device, list_devices
    if request.method == "GET":
        return jsonify(list_devices())
    body = request.get_json(force=True) or {}
    if request.method == "DELETE":
        ok = unregister_device(token=body.get("token"), device_id=body.get("device_id"))
        return jsonify({"ok": ok})
    device = register_device(
        body.get("token", ""),
        platform=body.get("platform", "web"),
        device_name=body.get("device_name", ""),
        company_filter=body.get("company_filter", ""),
    )
    return jsonify({"ok": True, "device": device})


@app.route("/api/notifications/mobile-config", methods=["GET", "PUT"])
def api_mobile_push_config():
    from risk_mobile_push import load_mobile_push_config, save_mobile_push_config
    if request.method == "GET":
        cfg = load_mobile_push_config()
        safe = dict(cfg)
        if safe.get("fcm_server_key"):
            safe["fcm_server_key"] = "******"
        return jsonify(safe)
    body = request.get_json(force=True) or {}
    body.pop("fcm_server_key", None)
    cfg = load_mobile_push_config()
    cfg.update(body)
    save_mobile_push_config(cfg)
    return jsonify({"ok": True})


@app.route("/api/actions/track", methods=["POST"])
def api_actions_track():
    try:
        body = request.get_json(force=True) or {}
        company = body.get("company_name", "")
        action_id = body.get("action_id", "")
        status = body.get("status", "completed")
        note = body.get("note", "")
        if not company or not action_id:
            return jsonify({"error": "缺少 company_name 或 action_id"}), 400
        from erm_auth import current_user
        from risk_workflow import apply_status, snapshot
        user = current_user()
        rec = apply_status(
            kind="action",
            item_id=action_id,
            company_name=company,
            status=status,
            user=user,
            reason=note or body.get("reason") or body.get("acceptance_reason") or "",
            review_date=body.get("review_date") or "",
            owner=body.get("owner") or "",
            due_date=body.get("due_date") or "",
            priority=body.get("priority") or "",
        )
        return jsonify({"ok": True, "action_id": action_id, "record": rec, "workflow": snapshot(company)})
    except PermissionError as exc:
        return jsonify({"error": str(exc), "code": "ERR-ERM-APPROVAL"}), 403
    except Exception as exc:
        return jsonify({"error": str(exc)}), 500


@app.route("/api/workflow", methods=["GET"])
def api_workflow_get():
    from erm_auth import current_user
    company = (request.args.get("company") or "").strip()
    if not company:
        return jsonify({"error": "缺少 company"}), 400
    if not current_user() and not company.startswith("DEMO-"):
        return jsonify({"error": "未登录", "code": "ERR-ERM-AUTH"}), 401
    from risk_workflow import snapshot
    return jsonify(snapshot(company))


@app.route("/api/workflow/status", methods=["POST"])
def api_workflow_status():
    try:
        body = request.get_json(force=True) or {}
        company = (body.get("company_name") or "").strip()
        item_id = (body.get("item_id") or body.get("id") or "").strip()
        kind = (body.get("kind") or "register").strip()
        status = body.get("status") or ""
        if not company or not item_id or not status:
            return jsonify({"error": "缺少 company_name / item_id / status"}), 400
        from erm_auth import current_user
        from risk_workflow import apply_status, snapshot
        rec = apply_status(
            kind=kind,
            item_id=item_id,
            company_name=company,
            status=status,
            user=current_user(),
            reason=body.get("reason") or body.get("acceptance_reason") or "",
            review_date=body.get("review_date") or "",
            owner=body.get("owner") or "",
            due_date=body.get("due_date") or "",
            priority=body.get("priority") or "",
            risk_rating=body.get("risk_rating") or "",
        )
        return jsonify({"ok": True, "record": rec, "workflow": snapshot(company)})
    except PermissionError as exc:
        return jsonify({"error": str(exc), "code": "ERR-ERM-APPROVAL"}), 403
    except Exception as exc:
        return jsonify({"error": str(exc)}), 500


@app.route("/api/bayesian", methods=["POST"])
def api_bayesian():
    try:
        data = request.get_json(force=True) or {}
        template = get_template_fields()
        if template:
            data = normalize_form_data(data.get("form_data", data), template)
        stats = compute_form_stats(data, template) if template else None
        result = run_assessment_from_form(data)
        from risk_bayesian import find_prior_assessment, run_bayesian_update
        prior = find_prior_assessment(load_history(), result.company_name)
        return jsonify(run_bayesian_update(result, prior, stats))
    except Exception as exc:
        return jsonify({"error": str(exc)}), 500


@app.route("/api/tasks", methods=["GET"])
def api_tasks_list():
    from erm_auth import current_user
    if not current_user():
        return jsonify({"error": "未登录", "code": "ERR-ERM-AUTH"}), 401
    from risk_tasks import list_tasks
    status = (request.args.get("status") or "open").strip()
    company = (request.args.get("company") or "").strip()
    rows = list_tasks(status="" if status == "all" else status, company=company)
    return jsonify({"tasks": rows, "total": len(rows)})


@app.route("/api/tasks/<task_id>/complete", methods=["POST"])
def api_tasks_complete(task_id):
    from erm_auth import current_user
    if not current_user() or not current_user().can_write():
        return jsonify({"error": "需要评估人或审批人", "code": "ERR-ERM-AUTH"}), 403
    from risk_tasks import complete_task
    rec = complete_task(task_id)
    if not rec:
        return jsonify({"error": "任务不存在"}), 404
    return jsonify({"ok": True, "task": rec})


@app.route("/api/knowledge/citations", methods=["GET"])
def api_knowledge_citations():
    from risk_knowledge import lookup_citations
    dim = request.args.get("dimension") or ""
    industry = request.args.get("industry") or ""
    cites = lookup_citations(dim, industry, limit=int(request.args.get("limit") or 3))
    return jsonify({"citations": cites})


if __name__ == "__main__":
    import threading
    import webbrowser

    from server_config import get_host, get_port, get_public_base_url, is_production

    if is_production() or os.environ.get("ERM_USE_WAITRESS") == "1":
        from run_production import main as run_prod
        run_prod()
        raise SystemExit(0)

    port = get_port()
    host = get_host()
    url = get_public_base_url()

    # 写入访问地址，便于用户查看
    try:
        with open(os.path.join(HISTORY_DIR, "server.url"), "w", encoding="utf-8") as f:
            f.write(url + "\n")
    except OSError:
        pass

    print("=" * 60)
    print("  企业风险动态评估系统 - Web 版 (开发模式)")
    print("=" * 60)
    print(f"  >>> 请在浏览器打开: {url}")
    print(f"  >>> 本地: http://127.0.0.1:{port}")
    print(f"  数据目录: {DATA_DIR}")
    tpl = _find_template_xlsx()
    if tpl:
        print(f"  Excel模板: {os.path.basename(tpl)}")
    else:
        print("  [警告] 未找到 Excel 模板，请运行 generate_form.py 生成")
    print("=" * 60)
    print("  生产部署请使用: py -3 run_production.py  或  生产环境启动.bat")
    print("=" * 60)

    def _open_browser():
        try:
            webbrowser.open(f"http://127.0.0.1:{port}")
        except Exception:
            pass

    threading.Timer(1.5, _open_browser).start()

    try:
        from bootstrap import init_server
        init_server()
        print("  定时 ERP 同步调度器已启动（间隔检查 60s）")
    except Exception as exc:
        print(f"  [警告] 调度器未启动: {exc}")

    try:
        app.run(debug=False, host=host, port=port, use_reloader=False)
    except OSError as e:
        print(f"\n[错误] 端口 {port} 无法启动: {e}")
        print("请关闭其他已运行的本系统窗口后重试，或设置环境变量 ERM_PORT=8090")
        input("\n按回车键退出...")
